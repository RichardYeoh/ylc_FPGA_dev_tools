# -*- coding: utf-8 -*-
"""Excel（.xlsx/.xlsm，以及纯 Python 读取的 .xls）转 Markdown。

要点：
- 每个工作表输出为二级标题 + 一张 md 表格；
- 合并单元格只保留左上角的值，其余留空；
- 默认读取公式缓存值；若整簿没有缓存值（文件从未被 Excel 保存过），
  自动退回读取公式原文并给出警告；
- 隐藏工作表默认跳过（可用选项包含）。
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import List, Sequence

from ..mdutil import render_table, squash_blank_lines, to_text

MAX_ROWS_WARN = 2000
MAX_COLS_WARN = 12


def _fmt_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        if isinstance(v, _dt.datetime):
            if v.hour == v.minute == v.second == 0:
                return v.strftime("%Y-%m-%d")
            return v.strftime("%Y-%m-%d %H:%M:%S")
        return v.isoformat()
    if isinstance(v, float):
        if v != v:
            return ""
        if v == int(v) and abs(v) < 1e16:
            return str(int(v))
        return ("%.10g" % v).rstrip("0").rstrip(".") or "0"
    return str(v).strip()


def _sheet_grid_xlsx(ws, opts: dict):
    """把 openpyxl 工作表转成二维字符串列表，并处理合并单元格。"""
    merged = {}
    for rng in ws.merged_cells.ranges:
        try:
            top_left = ws.cell(row=rng.min_row, column=rng.min_col).value
        except Exception:
            top_left = None
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = top_left if (r == rng.min_row and c == rng.min_col) else None

    grid: List[List[str]] = []
    for row in ws.iter_rows():
        values = []
        for cell in row:
            key = (cell.row, cell.column)
            v = merged[key] if key in merged else cell.value
            values.append(_fmt_value(v))
        grid.append(values)
    return grid


def _count_missing_formula_values(src: Path, wb_val) -> int:
    """统计"有公式但没有缓存值"的单元格数量（Excel 未保存过时常见）。"""
    from openpyxl import load_workbook

    try:
        wb_f = load_workbook(str(src), data_only=False, read_only=True)
    except Exception:
        return 0
    count = 0
    for ws in wb_f.worksheets:
        ws_val = wb_val[ws.title] if ws.title in wb_val.sheetnames else None
        if ws_val is None:
            continue
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    try:
                        if ws_val[cell.coordinate].value is None:
                            count += 1
                    except Exception:
                        continue
    try:
        wb_f.close()
    except Exception:
        pass
    return count


def _trim(grid: List[List[str]]):
    """去掉全空行与全空列，保留相对顺序。"""
    rows = [r for r in grid if any((c or "").strip() for c in r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [i for i in range(width) if any(r[i].strip() for r in rows)]
    return [[r[i] for i in keep] for r in rows]


def convert_xlsx(src: Path, md_path: Path, options: dict | None = None) -> tuple[str, List[str]]:
    from openpyxl import load_workbook

    opt = options or {}
    include_hidden = bool(opt.get("include_hidden_sheets", False))
    warnings: List[str] = []

    wb = load_workbook(str(src), data_only=True, read_only=False)
    has_values = any(
        v is not None for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row
    )
    if not has_values:
        wb = load_workbook(str(src), data_only=False, read_only=False)
        warnings.append(
            "工作簿没有公式缓存值，公式按原文（如 =SUM(A1:A3)）输出；建议先用 Excel 打开并保存一次"
        )
    else:
        missing = _count_missing_formula_values(src, wb)
        if missing:
            warnings.append(
                f"{missing} 个公式单元格没有缓存值（输出为空）；如需要请用 Excel 打开并保存一次"
            )

    blocks: List[str] = []
    for ws in wb.worksheets:
        if not include_hidden and getattr(ws, "sheet_state", "visible") != "visible":
            continue
        grid = _trim(_sheet_grid_xlsx(ws, opt))
        if not grid:
            blocks.append(f"## {ws.title}\n\n（空工作表）")
            continue
        if len(grid) > MAX_ROWS_WARN:
            warnings.append(f"工作表「{ws.title}」有 {len(grid)} 行，已全部输出，md 文件会较长")
        if grid and len(grid[0]) > MAX_COLS_WARN:
            warnings.append(f"工作表「{ws.title}」有 {len(grid[0])} 列，md 表格显示会比较拥挤")
        table = render_table(grid, has_header=True)
        blocks.append(f"## {ws.title}\n\n{table}")

    md = f"# {opt.get('title_name') or Path(src).stem}\n\n" + "\n\n".join(blocks)
    return squash_blank_lines(md).strip() + "\n", warnings


def convert_xls(src: Path, md_path: Path, options: dict | None = None) -> tuple[str, List[str]]:
    """老格式 .xls：用 xlrd 纯 Python 读取，不依赖 Office。"""
    import xlrd

    opt = options or {}
    warnings: List[str] = []
    include_hidden = bool(opt.get("include_hidden_sheets", False))

    try:
        book = xlrd.open_workbook(str(src), formatting_info=True)
    except Exception as exc:
        book = xlrd.open_workbook(str(src), formatting_info=False)
        warnings.append(f"无法读取单元格合并信息，按普通网格输出：{exc}")

    blocks: List[str] = []
    for idx in range(book.nsheets):
        sheet = book.sheet_by_index(idx)
        if not include_hidden and getattr(sheet, "visibility", 0) != 0:
            continue
        grid = _xls_grid(sheet, book)
        grid = _trim(grid)
        if not grid:
            blocks.append(f"## {sheet.name}\n\n（空工作表）")
            continue
        if len(grid) > MAX_ROWS_WARN:
            warnings.append(f"工作表「{sheet.name}」有 {len(grid)} 行，md 文件会较长")
        blocks.append(f"## {sheet.name}\n\n{render_table(grid, has_header=True)}")

    md = f"# {opt.get('title_name') or Path(src).stem}\n\n" + "\n\n".join(blocks)
    return squash_blank_lines(md).strip() + "\n", warnings


def _xls_grid(sheet, book) -> List[List[str]]:
    """把 xlrd 工作表转二维字符串，合并区域只保留左上角值。"""
    merged_map: dict[tuple[int, int], object] = {}
    for rlo, rhi, clo, chi in getattr(sheet, "merged_cells", []):
        top_left = sheet.cell_value(rlo, clo) if rlo < sheet.nrows and clo < sheet.ncols else ""
        for r in range(rlo, min(rhi, sheet.nrows)):
            for c in range(clo, min(chi, sheet.ncols)):
                merged_map[(r, c)] = top_left if (r == rlo and c == clo) else None

    grid: List[List[str]] = []
    for r in range(sheet.nrows):
        row: List[str] = []
        for c in range(sheet.ncols):
            if (r, c) in merged_map:
                value = merged_map[(r, c)]
                row.append("" if value is None else _fmt_value(value))
            else:
                row.append(_fmt_xlrd_cell(sheet.cell(r, c), book))
        grid.append(row)
    return grid


def _fmt_xlrd_cell(cell, book) -> str:
    import xlrd

    ctype = cell.ctype
    if ctype == xlrd.XL_CELL_EMPTY or ctype == xlrd.XL_CELL_BLANK:
        return ""
    if ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if ctype == xlrd.XL_CELL_DATE:
        try:
            d = xlrd.xldate_as_datetime(cell.value, book.datemode)
            if d.hour == d.minute == d.second == 0:
                return d.strftime("%Y-%m-%d")
            return d.strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            return to_text(cell.value)
    if ctype == xlrd.XL_CELL_ERROR:
        return "#ERR"
    if ctype == xlrd.XL_CELL_NUMBER:
        return _fmt_value(float(cell.value))
    return str(cell.value).strip()
