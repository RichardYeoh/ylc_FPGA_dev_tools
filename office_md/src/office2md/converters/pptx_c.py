# -*- coding: utf-8 -*-
"""PowerPoint（.pptx/.pptm）转 Markdown。

映射规则：
- 每张幻灯片 → `## <序号>. <标题>`；
- 文本占位符/文本框 → 普通段落，级别 > 0 的段落用缩进列表；
- 图片 → 导出到 `<md名>_files/` 并内联引用；
- 表格 → md 表格；图表 → 由图例/类别/数值还原成 md 表格；
- 演讲者备注 → 引用块（可用选项关闭）。
"""

from __future__ import annotations

from pathlib import Path
from typing import List

from ..mdutil import ImageStore, render_table, squash_blank_lines

_LEVEL_INDENT = "  "


def convert_pptx(src: Path, md_path: Path, options: dict | None = None) -> tuple[str, List[str]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    opt = options or {}
    with_notes = bool(opt.get("include_notes", True))
    images = ImageStore(md_path)
    warnings: List[str] = []

    prs = Presentation(str(src))
    blocks: List[str] = []
    blocks.append(f"# {opt.get('title_name') or Path(src).stem}")

    for idx, slide in enumerate(prs.slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip().replace("\n", " ")
        except Exception:
            title = ""
        head = f"## {idx}. {title}" if title else f"## {idx}. 第{idx}页"
        parts: List[str] = [head]

        for shape in slide.shapes:
            try:
                chunk = _render_shape(shape, images, warnings, MSO_SHAPE_TYPE, skip_title=bool(title))
            except Exception as exc:  # 单个形状失败不影响整页
                warnings.append(f"第 {idx} 页存在无法解析的形状（{type(shape).__name__}）：{exc}")
                chunk = ""
            if chunk:
                parts.append(chunk)

        if with_notes:
            note = _notes_text(slide)
            if note:
                quoted = "\n".join("> " + line for line in note.split("\n"))
                parts.append("> **备注**\n>\n" + quoted)

        blocks.append("\n\n".join(parts))

    md = "\n\n".join(blocks)
    return squash_blank_lines(md).strip() + "\n", warnings


def _notes_text(slide) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        tf = slide.notes_slide.notes_text_frame
        if tf is None:
            return ""
        text = (tf.text or "").strip()
    except Exception:
        return ""
    # 去掉纯数字页码
    if text.isdigit():
        return ""
    return text


def _render_shape(shape, images: ImageStore, warnings: List[str], MSO_SHAPE_TYPE, skip_title: bool = False) -> str:
    st = getattr(shape, "shape_type", None)

    if st == MSO_SHAPE_TYPE.GROUP:
        out = []
        for sub in shape.shapes:
            chunk = _render_shape(sub, images, warnings, MSO_SHAPE_TYPE)
            if chunk:
                out.append(chunk)
        return "\n\n".join(out)

    # 表格
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = []
        for row in shape.table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        return render_table(rows, has_header=True)

    # 图表
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return _chart_table(shape.chart, warnings)

    # 图片
    if st == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
        try:
            img = shape.image
            md = images.markdown(img.blob, "." + (img.ext or "png"))
            if md:
                return md
        except Exception:
            pass

    # 文本
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        if skip_title and _is_title_placeholder(shape):
            return ""
        return _text_frame(shape.text_frame)

    if st == MSO_SHAPE_TYPE.LINE:
        return ""
    return ""


def _is_title_placeholder(shape) -> bool:
    try:
        ph = shape.placeholder_format
        return ph is not None and ph.idx == 0
    except Exception:
        return False


def _text_frame(tf) -> str:
    out: List[str] = []
    prev_is_list = False
    for para in tf.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        level = int(getattr(para, "level", 0) or 0)
        if level > 0:
            item = _LEVEL_INDENT * (level - 1) + "- " + text
            cur_is_list = True
        else:
            item = text
            cur_is_list = False
        if out and cur_is_list and prev_is_list:
            out[-1] = out[-1] + "\n" + item
        else:
            out.append(item)
        prev_is_list = cur_is_list
    return "\n\n".join(out)


def _chart_table(chart, warnings: List[str]) -> str:
    try:
        cats: List[str] = []
        series_rows: List[List[str]] = []
        names: List[str] = []
        for plot in chart.plots:
            try:
                cats = [str(c) for c in (plot.categories or [])]
            except Exception:
                cats = []
            for s in plot.series:
                names.append(str(s.name) if s.name else "系列")
                vals = list(s.values or [])
                series_rows.append(["" if v is None else str(v) for v in vals])
        if not series_rows:
            return ""
        nrow = max([len(r) for r in series_rows] + [len(cats)])
        header = ["类别"] + names
        rows = [header]
        for i in range(nrow):
            cat = cats[i] if i < len(cats) else str(i + 1)
            rows.append([cat] + [r[i] if i < len(r) else "" for r in series_rows])
        title = ""
        try:
            if chart.has_title and chart.chart_title is not None:
                title = (chart.chart_title.text_frame.text or "").strip()
        except Exception:
            title = ""
        table = render_table(rows, has_header=True)
        return (f"**图表：{title}**\n\n{table}" if title else table)
    except Exception as exc:
        warnings.append(f"图表解析失败：{exc}")
        return ""
